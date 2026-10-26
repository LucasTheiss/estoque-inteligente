"""Completa o XLSX coletivo existente, preservando as fórmulas e o modelo original.
Dependências: openpyxl e Pillow. Recalcular o arquivo produzido no LibreOffice/Excel.
"""
from pathlib import Path
from copy import copy
import json, math, re, hashlib, shutil, textwrap
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill, Border
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.page import PageMargins
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'docs/apf/coletiva'
target=BASE/'entrega-coletiva-apf.xlsx'
w=openpyxl.load_workbook(target)
data=json.loads((BASE/'catalogo-revisado.json').read_text(encoding='utf-8'))
functions={x['id']:x for x in data['dados']+data['funcoes']}
manifest={}
for name in ['manifesto.json','manifesto-final.json']:
 for m in json.loads((BASE/'capturas'/name).read_text(encoding='utf-8')):
  path=BASE/'capturas'/Path(m['arquivo']).name
  assert path.is_file(),path
  m['arquivo']=str(path.relative_to(ROOT)).replace('\\','/')
  manifest[m['id']]=m
for m in manifest.values():
 assert hashlib.sha256((ROOT/m['arquivo']).read_bytes()).hexdigest()==m['sha256'],m['arquivo']
(BASE/'capturas/manifesto.json').write_text(json.dumps(list(manifest.values()),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def setval(sheet,cell,value):w[sheet][cell]=value
setval('Capa','B7','Branch entrega-trabalho — conclusão da revisão de 25/10/2026')
setval('Capa','B15','Notas NI são avaliação técnica com premissas acadêmicas explícitas. Capturas reais incluem dados de demonstração e simulador IoT identificado. Os três recortes individuais são documentos históricos preservados.')
setval('Critérios e fontes','B31','Capturas reais incorporadas em todas as 55 fichas, com origem e legenda. As EEs incluem entrada e resultado; cada CE possui entrada e saída. Balança e LED usam simulador MQTT identificado. Ver capturas/manifesto-incorporado.json.')
setval('Critérios e fontes','B33','Suíte registrada na revisão anterior de 25/10/2026: 90 testes Java e 13 JavaScript aprovados, conforme logs preservados. Na conclusão foram verificadas capturas, fórmulas, recálculo e impressão; não se alega nova execução dessa suíte.')
setval('Critérios e fontes','B35','Integrantes: Artur Ribeiro de Souza, Cristhian Ruan de Rezende e Lucas Andriani Theiss. Identificação e classificação apresentadas como análise coletiva; não há declaração inventada de autoria individual ou de consenso das notas.')
setval('Critérios e fontes','B27','Contagem coletiva atual sincronizada em docs/apf/contagem.json, resultados.json, CALCULOS.md, FICHAS.md e NI.md. Recortes individuais e PDF de divisão preservam a revisão anterior e não devem ser somados.')
setval('Conferência','B15','14 notas justificadas em NI e estimativas. Volume e múltiplos locais são premissas acadêmicas; a pontuação é avaliação técnica documental.')
setval('Conferência','B11','55 fichas com 93 imagens reais incorporadas. Todas as EEs e CEs possuem entrada e saída; balança e LED mostram resposta real do simulador identificado, e SE09 inclui o PDF efetivamente exportado.')
setval('Conferência','B31','Fontes coletivas sincronizadas com o catálogo revisado; os recortes individuais e PDF de divisão permanecem históricos. Consulte o Excel coletivo para os totais atuais.')
setval('Conferência','B33','Todas as fichas incluem capturas reais correspondentes. Legendas distinguem seleção, formulário, resultado e simulador. A captura de entrada não é apresentada como confirmação de operação.')
setval('Conferência','B35','Convenções de contagem e premissas estão explícitas. Não se declara consenso da equipe nem autoria individual não informada. Validação de compatibilidade realizada no LibreOffice; Microsoft Excel indisponível nesta máquina.')
setval('SE08','B5','/movimentar (loja granel); GET /api/iot/balanca/{produto}/{device}')
setval('SE08','B7','IoTService.peso; movimento.html; app.js; DispositivosDemo. Leitura recente do simulador MQTT capturada na conclusão.')
setval('SE08','B8','A02: produto em kg e preço. A08: loja e configuração MQTT. Bruto, tara e indicador de simulação chegam do dispositivo; líquido e subtotal são calculados e devolvidos pela API. A tela mostra líquido, subtotal e simulação.')
setval('Exemplos SE','B9',1.440)
setval('Exemplos SE','C9',0.015)
setval('Exemplos SE','C10',29.90)
setval('Exemplos SE','B22','Captura real na loja granel: 1,440 kg bruto − 0,015 kg tara = 1,425 kg líquido; × R$ 29,90/kg = R$ 42,61. Simulador MQTT identificado, leitura com até 5 segundos; nenhum movimento registrado.')
w.calculation=CalcProperties(calcId=191029,calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
ink='243746';blue='DCEAF4'
incorporated=[]
for key,f in functions.items():
 s=w[key]
 s._images=[]
 for column in ['C','D']:
  seed=s[f'{column}11'].value
  for r in range(11,43):
   s[f'{column}{r}']=Translator(seed,origin=f'{column}11').translate_formula(f'{column}{r}')
 if f['tipo']=='CE':
  s['B57']=f['regraAR']
  for r in range(46,57):
   ref=s.cell(r,1).value
   if ref:
    s.cell(r,3,int(ref in f.get('arquivosEntrada',f['arquivos'])))
    s.cell(r,4,int(ref in f.get('arquivosSaida',f['arquivos'])))
 # Ajusta alturas sem mover fórmulas ou mudar a organização das fichas.
 for r in range(1,75):
  s.row_dimensions[r].hidden=False
  s.row_dimensions[r].height=22
 for r in [9,43,58,66,68,74]:s.row_dimensions[r].height=8
 for r in range(11,43):
  if s.cell(r,1).value is None:
   s.row_dimensions[r].hidden=True
  else:s.row_dimensions[r].height=30 if len(str(s.cell(r,1).value))>27 else 21
 for r in range(46,57):
  if s.cell(r,1).value is None:s.row_dimensions[r].hidden=True
  else:s.row_dimensions[r].height=34
 for r in [1,10,44,45,59,69]:s.row_dimensions[r].height=28
 for r in [2,3,4,5,6,7,8,57,67,70,71,72,73]:
  value=str(s.cell(r,2).value or '')
  lines=sum(max(1,math.ceil(len(x)/100)) for x in value.split('\n'))
  s.row_dimensions[r].height=max(29,lines*15+10)
 s['A69']='Fluxo da função — leitura do protótipo real'
 s['A70']='Entrada / contexto'
 s['A72']='Resultado / saída'
 s['A73']='Evidência visual'
 s['B73']='Capturas da aplicação nas páginas seguintes. Veja a legenda de cada estado e a origem registrada.'
 s.row_dimensions[73].height=36
 # Todas as imagens antigas são substituídas por fontes identificadas no manifesto.
 for merged in list(s.merged_cells.ranges):
  if merged.min_row>=75:s.unmerge_cells(str(merged))
 for row in s.iter_rows(min_row=75,max_row=max(190,s.max_row),max_col=4):
  for c in row:
   c.value=None;c._style=copy(s['B2']._style);c.border=Border();c.fill=PatternFill()
 for r in range(75,191):s.row_dimensions[r].height=20;s.row_dimensions[r].hidden=False
 if f['tipo'] in ['EE','CE']:
  ids=[key+'-entrada',key+'-saida']
 elif key in ['SE08','SE09']:ids=[key+'-entrada',key+'-saida']
 else:ids=[key]
 assert all(i in manifest for i in ids),(key,ids)
 s.row_breaks=RowBreak()
 for r in [43,74]:s.row_breaks.append(Break(id=r))
 for index,capid in enumerate(ids):
  m=manifest[capid];path=ROOT/m['arquivo'];r=75+56*index
  if index:s.row_breaks.append(Break(id=r-1))
  for rr in [r,r+1,r+2]:s.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=4)
  s.cell(r,1,f'{key} — captura real {index+1}/{len(ids)}')
  s.cell(r,1).font=Font(name='Arial',size=13,bold=True,color=ink)
  s.cell(r,1).fill=PatternFill('solid',fgColor=blue)
  s.cell(r+1,1,m['descricao']);s.row_dimensions[r+1].height=42
  s.cell(r+2,1,f"Origem: {m['url']} | Capturada em: {m['data']} | {Path(m['arquivo']).name}")
  s.row_dimensions[r+2].height=30
  img=Image(path);ratio=min(980/img.width,1280/img.height)
  img.width*=ratio;img.height*=ratio;s.add_image(img,f'A{r+4}')
  incorporated.append(dict(funcao=key,**m))
 last=75+56*len(ids)-1
 s.print_area=f'A1:D{last}'
 s.print_options.horizontalCentered=True
 s.sheet_view.zoomScale=85

for s in w:
 s.page_setup.orientation='portrait' if s.title in functions or s.title in ['NI e estimativas','Exemplos SE'] else 'landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
 s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
 s.sheet_properties.pageSetUpPr.fitToPage=True
 s.page_margins=PageMargins(left=.25,right=.25,top=.35,bottom=.35,header=.15,footer=.15)
 s.oddFooter.center.text='Estoque Inteligente — &A — página &P'
 s.oddFooter.center.size=8
 for row in s:
  for c in row:
   if c.value is not None:
    a=copy(c.alignment);a.wrap_text=True;a.vertical='center';c.alignment=a
 if s.title not in functions:
  populated=[c for row in s for c in row if c.value is not None]
  last=max(c.row for c in populated);col=max(4,max(c.column for c in populated))
  s.print_area=f'A1:{openpyxl.utils.get_column_letter(col)}{last}'
  for r in range(1,last+1):
   height=22
   for c in s[r]:
    if c.value is None:continue
    if c.data_type=='f':continue
    width=s.column_dimensions[c.column_letter].width or 13
    for mr in s.merged_cells.ranges:
     if c.coordinate in mr:
      width=sum(s.column_dimensions[openpyxl.utils.get_column_letter(i)].width or 13 for i in range(mr.min_col,mr.max_col+1));break
    lines=sum(max(1,math.ceil(len(line)/max(6,width-4))) for line in str(c.value).split('\n'))
    height=max(height,lines*16+12)
   s.row_dimensions[r].height=height
  if s.title=='Catálogo':s.print_title_rows='1:3'
  if s.title=='NI e estimativas':
   s.row_breaks=RowBreak();s.row_breaks.append(Break(id=18))
 # Número e rótulo seguinte ficam em células distintas, com alinhamento definido.
 for row in s:
  for c in row:
   if c.data_type=='f' and not str(c.value).startswith('=HYPERLINK'):
    a=copy(c.alignment);a.horizontal='center';c.alignment=a

ni=w['NI e estimativas']
ni.data_validations.dataValidation=[]
dv=DataValidation(type='whole',operator='between',formula1=0,formula2=5);dv.errorTitle='Nota inválida';dv.error='Informe um inteiro de 0 a 5.';dv.showErrorMessage=True
ni.add_data_validation(dv);dv.add('B4:B17')
for cell in ['B27','B28','B29','B30','B32']:
 dv=DataValidation(type='decimal',operator='greaterThan',formula1=0);dv.showErrorMessage=True;dv.error='Informe um valor maior que zero.';ni.add_data_validation(dv);dv.add(cell)
for cell in ['B31','B33','B34']:
 dv=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0);dv.showErrorMessage=True;dv.error='Informe um valor não negativo.';ni.add_data_validation(dv);dv.add(cell)
w.save(target)
(BASE/'capturas/manifesto-incorporado.json').write_text(json.dumps(incorporated,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('XLSX atualizado:',len(w.sheetnames),'abas;',len(incorporated),'imagens incorporadas; recalcular antes de entregar.')
