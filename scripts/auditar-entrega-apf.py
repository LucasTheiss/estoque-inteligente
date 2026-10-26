"""Auditoria de preservação, caches, imagens e impressão da entrega recalculada."""
from pathlib import Path
import json,hashlib,re,math
import openpyxl,pymupdf
from PIL import Image,ImageDraw
root=Path(__file__).resolve().parents[1];base=root/'docs/apf/coletiva'
originals=json.loads((base/'originais-sha256.json').read_text())
for name,sha in originals.items():assert hashlib.sha256((root/name.replace('\\','/')).read_bytes()).hexdigest()==sha,name
w=openpyxl.load_workbook(base/'entrega-coletiva-apf.xlsx',data_only=False)
c=openpyxl.load_workbook(base/'entrega-coletiva-apf.xlsx',data_only=True)
formulas=[(s.title,cell.coordinate) for s in w for row in s for cell in row if cell.data_type=='f']
errors=[(s.title,cell.coordinate,cell.value) for s in c for row in s for cell in row if cell.data_type=='e']
assert not errors,errors
missing=[(s,r) for s,r in formulas if c[s][r].value is None];assert not missing,missing
manifest=json.loads((base/'capturas/manifesto-incorporado.json').read_text())
for m in manifest:assert hashlib.sha256((root/m['arquivo']).read_bytes()).hexdigest()==m['sha256'],m['arquivo']
assert sum(len(s._images) for s in w)==len(manifest)
doc=pymupdf.open(root/'tmp/apf-final-impressao.pdf')
images=[];outside=[];blank=[];pages=[]
for i,p in enumerate(doc):
 info=p.get_image_info()
 for item in info:
  rect=pymupdf.Rect(item['bbox']);images.append(dict(pagina=i+1,bbox=list(rect)))
  if rect.x0<-.1 or rect.y0<-.1 or rect.x1>p.rect.width+.1 or rect.y1>p.rect.height+.1:outside.append(i+1)
 text=p.get_text();pages.append(text)
 if len(text.strip())<65 and not info:blank.append(i+1)
assert not outside,('Imagens fora da página',outside)
assert len(images)==len(manifest),('Número de imagens na impressão',len(images),len(manifest))
def norm(t):return re.sub(r'\s+','',str(t)).casefold()
by_sheet={s.title:norm('\n'.join(t for t in pages if norm('Estoque Inteligente — '+s.title+' — página') in norm(t))) for s in w}
assert all(by_sheet.values()),('Abas sem páginas identificadas',[k for k,v in by_sheet.items() if not v])
unmatched=[]
for s in w:
 if not s.print_area:continue
 for row in s:
  for cell in row:
   if not isinstance(cell.value,str) or cell.data_type=='f' or len(cell.value)<15 or s.row_dimensions[cell.row].hidden:continue
   if norm(cell.value) not in by_sheet[s.title]:unmatched.append(dict(aba=s.title,celula=cell.coordinate,texto=cell.value))
out=root/'tmp/apf-previews';out.mkdir(exist_ok=True)
for start in range(0,len(doc),24):
 canvas=Image.new('RGB',(1200,1800),'#ddd');draw=ImageDraw.Draw(canvas)
 for n in range(start,min(start+24,len(doc))):
  p=doc[n];pix=p.get_pixmap(matrix=pymupdf.Matrix(.23,.23),alpha=False);im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
  im.thumbnail((285,270));x=(n-start)%4*300;y=(n-start)//4*300
  canvas.paste(im,(x,y+20));draw.text((x+4,y+3),str(n+1),fill='black')
 canvas.save(out/f'contato-{start//24+1:02}.jpg')
report=dict(abas=len(w.sheetnames),formulas=len(formulas),erros=errors,cachesAusentes=missing,originaisPreservados=len(originals),imagensIncorporadas=len(manifest),paginasImpressao=len(doc),imagensNaImpressao=len(images),imagensForaDaPagina=outside,paginasSemConteudo=blank,textosNaoLocalizados=unmatched)
(base/'validacao-impressao.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='textosNaoLocalizados'},ensure_ascii=False));print('Textos não localizados:',len(unmatched));print(unmatched[:6])
